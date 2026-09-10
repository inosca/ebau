import datetime
import logging
import signal
import sqlite3
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import openpyxl
import pytest
from django.core.management import CommandError
from django.db import transaction

from camac.instance.management.commands import identify_and_fix_calculations
from camac.instance.management.commands.identify_and_fix_calculations import (
    Command,
    RecalcInfo,
    RunState,
    RunStateStore,
    StatsEntry,
    TaskResult,
)

CONFIGURATION = {
    "do_validate": False,
    "do_recalc": True,
    "do_fix_calc": True,
    "do_commit": True,
    "analyze_only": (),
    "search_only": None,
}


def _stats_entry(instance_id, event="detected"):
    return StatsEntry(
        "Test municipality",
        f"{instance_id} (in progress)",
        [
            RecalcInfo(
                name="calculated-field",
                hidden=False,
                before=1,
                after=2,
                event=event,
            )
        ],
    )


def _checkpoint_and_raise(state_configuration, result, message):
    with transaction.atomic():
        Command._checkpoint_on_commit(state_configuration, result)
        raise RuntimeError(message)


def test_run_state_is_created_resumed_and_reset():
    store = RunStateStore(CONFIGURATION)

    state = store.start()
    assert store.path.exists()
    assert store.path.name == "state.sqlite3"
    assert state.configuration == CONFIGURATION

    store.record_completed("recalculation", 1, [_stats_entry(1)])
    store.record_completed("recalculation", 1, [_stats_entry(1, event="fixed")])
    store.record_completed("validation", 2, [])
    resumed = store.start()
    assert resumed.completed_ids == {"recalculation": {1}, "validation": {2}}
    assert resumed.results[("recalculation", 1)] == [_stats_entry(1)]
    assert resumed.results[("validation", 2)] == []

    incompatible_store = RunStateStore({**CONFIGURATION, "do_fix_calc": False})
    with pytest.raises(CommandError, match="different command options"):
        incompatible_store.start()

    reset = incompatible_store.start(reset=True)
    assert reset.configuration == incompatible_store.configuration
    assert reset.completed_ids == {}


def test_checkpoint_results_are_stored_as_individual_sqlite_rows():
    store = RunStateStore(CONFIGURATION)
    store.start()
    store.record_completed("validation", 1, [])
    store.record_completed("recalculation", 1, [_stats_entry(1)])

    with sqlite3.connect(store.path) as connection:
        rows = connection.execute(
            "SELECT phase, instance_id FROM completed_instances ORDER BY phase"
        ).fetchall()

    assert rows == [("recalculation", 1), ("validation", 1)]


def test_malformed_sqlite_checkpoint_can_be_reset():
    store = RunStateStore(CONFIGURATION)
    store.directory.mkdir(parents=True)
    store.path.write_text("not an SQLite database")

    with pytest.raises(CommandError, match="Use --reset-state"):
        store.start()

    assert store.start(reset=True).configuration == CONFIGURATION


def test_many_sqlite_checkpoints_restore_all_completed_instances():
    store = RunStateStore(CONFIGURATION)
    store.start()

    for instance_id in range(1_000):
        store.record_completed("validation", instance_id, [])

    resumed = store.start()
    assert resumed.completed_ids == {"validation": set(range(1_000))}
    assert len(resumed.results) == 1_000


def test_run_configuration_ignores_verbosity_and_dev_mode():
    options = {**CONFIGURATION, "dev_mode": False, "verbosity": 0}

    assert Command._run_configuration(options) == CONFIGURATION
    assert (
        Command._run_configuration({**options, "dev_mode": True, "verbosity": 3})
        == CONFIGURATION
    )


def test_parallel_argument_accepts_an_optional_worker_count():
    parser = Command.__new__(Command).create_parser(
        "manage.py", "identify_and_fix_calculations"
    )

    assert parser.parse_args([]).parallel is False
    assert parser.parse_args(["--parallel"]).parallel is None
    assert parser.parse_args(["--parallel", "2"]).parallel == 2


@pytest.mark.parametrize("parallel", [None, 2])
def test_parallel_worker_count_is_passed_to_executor(monkeypatch, parallel):
    class RecordingExecutor:
        instance = None

        def __init__(self, max_workers=None):
            type(self).instance = self
            self.max_workers = max_workers

        def shutdown(self, *args, **kwargs):
            pass

    class StateStoreStub:
        def __init__(self, configuration):
            self.configuration = configuration

        def start(self, reset=False):
            return RunState(
                configuration=self.configuration,
                started_at=datetime.datetime.now(),
            )

    command = Command.__new__(Command)
    command.configure_logging = lambda verbosity: None
    command._restore_stats = lambda: None
    command.log_stats = lambda: None
    monkeypatch.setattr(
        identify_and_fix_calculations, "ProcessPoolExecutor", RecordingExecutor
    )
    monkeypatch.setattr(identify_and_fix_calculations, "RunStateStore", StateStoreStub)

    command.handle(
        verbosity=0,
        parallel=parallel,
        do_validate=False,
        do_recalc=False,
        do_fix_calc=False,
        do_commit=False,
        excel_out=None,
        dev_mode=False,
        analyze_only=None,
        search_only=None,
        reset_state=False,
    )

    assert RecordingExecutor.instance.max_workers == parallel


@pytest.mark.django_db(transaction=True)
def test_checkpoint_is_written_only_after_document_commit():
    store = RunStateStore(CONFIGURATION)
    store.start()
    assert store.start().completed_ids == {}
    result = TaskResult("recalculation", 1, [_stats_entry(1)])

    with pytest.raises(RuntimeError, match="rollback"):
        _checkpoint_and_raise(CONFIGURATION, result, "rollback")

    assert store.start().completed_ids == {}

    with transaction.atomic():
        Command._checkpoint_on_commit(CONFIGURATION, result)

    assert store.start().completed_ids == {"recalculation": {1}}


@pytest.mark.django_db(transaction=True)
def test_checkpoint_keeps_previously_committed_documents_after_failure():
    store = RunStateStore(CONFIGURATION)
    store.start()

    with transaction.atomic():
        Command._checkpoint_on_commit(
            CONFIGURATION, TaskResult("recalculation", 1, [_stats_entry(1)])
        )

    with pytest.raises(RuntimeError, match="second document"):
        _checkpoint_and_raise(
            CONFIGURATION,
            TaskResult("recalculation", 2, [_stats_entry(2)]),
            "second document failed",
        )

    state = store.start()
    assert state.completed_ids == {"recalculation": {1}}
    assert ("recalculation", 2) not in state.results


@pytest.mark.django_db(transaction=True)
@pytest.mark.parametrize(
    ("task", "task_kwargs", "phase", "operation"),
    [
        (Command.check_validation, {}, "validation", "validation"),
        (
            Command.check_miscalculation,
            {"analyze_only": None, "fix_calculation": False, "do_commit": False},
            "recalculation",
            "recalculation",
        ),
    ],
)
def test_failed_tasks_are_logged_without_being_checkpointed(
    monkeypatch, caplog, task, task_kwargs, phase, operation
):
    def fail_clean_env(cls):
        raise RuntimeError("expected failure")

    RunStateStore(CONFIGURATION).start(reset=True)
    monkeypatch.setattr(Command, "_clean_env", classmethod(fail_clean_env))

    result = task("[Test 1/1]", 1, state_configuration=CONFIGURATION, **task_kwargs)

    command = Command.__new__(Command)
    command.stats = defaultdict(lambda: defaultdict(list))
    with caplog.at_level(logging.ERROR):
        command._log_result(SimpleNamespace(result=lambda: result))

    assert result.phase == phase
    assert any(
        f"Error during {operation} (expected failure)" in message
        for message in caplog.messages
    )
    assert RunStateStore(CONFIGURATION).start().completed_ids == {}


def test_concurrent_checkpoint_merges_preserve_all_outcomes():
    RunStateStore(CONFIGURATION).start()

    def record(instance_id):
        RunStateStore(CONFIGURATION).record_completed(
            "recalculation", instance_id, [_stats_entry(instance_id)]
        )

    with ThreadPoolExecutor(max_workers=4) as executor:
        list(executor.map(record, range(1, 9)))

    state = RunStateStore(CONFIGURATION).start()
    assert state.completed_ids == {"recalculation": set(range(1, 9))}
    assert set(state.results) == {
        ("recalculation", instance_id) for instance_id in range(1, 9)
    }


@pytest.mark.parametrize("stop_signal", [signal.SIGINT, signal.SIGTERM])
def test_stop_signals_cancel_pending_work_without_waiting(monkeypatch, stop_signal):
    class InterruptingExecutor:
        instance = None

        def __init__(self):
            type(self).instance = self
            self.shutdown_calls = []

        def submit(self, *args, **kwargs):
            signal.raise_signal(stop_signal)

        def shutdown(self, *args, **kwargs):
            self.shutdown_calls.append((args, kwargs))

    class StateStoreStub:
        def __init__(self, configuration):
            self.configuration = configuration

        def start(self, reset=False):
            return RunState(
                configuration=self.configuration,
                started_at=datetime.datetime.now(),
            )

    command = Command.__new__(Command)
    command.configure_logging = lambda verbosity: None
    command._restore_stats = lambda: None
    command.iter_instances_for_validation = lambda: (1, [SimpleNamespace(pk=1)])
    command.log_stats = lambda: pytest.fail("stats should not be logged on interrupt")
    previous_handler = signal.getsignal(stop_signal)

    monkeypatch.setattr(
        identify_and_fix_calculations, "SynchronousExecutor", InterruptingExecutor
    )
    monkeypatch.setattr(identify_and_fix_calculations, "RunStateStore", StateStoreStub)

    command.handle(
        verbosity=0,
        parallel=False,
        do_validate=True,
        do_recalc=False,
        do_fix_calc=False,
        do_commit=False,
        excel_out=None,
        dev_mode=False,
        analyze_only=None,
        search_only=None,
        reset_state=False,
    )

    assert signal.getsignal(stop_signal) is previous_handler
    assert InterruptingExecutor.instance.shutdown_calls == [
        ((), {"wait": False, "cancel_futures": True})
    ]


class QuerySetStub:
    def __init__(self, instance_ids):
        self.instance_ids = list(instance_ids)

    def exclude(self, **kwargs):
        excluded_ids = kwargs.get("pk__in")
        if excluded_ids is None:
            return self
        return QuerySetStub(
            instance_id
            for instance_id in self.instance_ids
            if instance_id not in excluded_ids
        )

    def filter(self, **kwargs):
        included_ids = kwargs["pk__in"]
        return QuerySetStub(
            instance_id
            for instance_id in self.instance_ids
            if instance_id in included_ids
        )

    def order_by(self, *args):
        return self

    def count(self):
        return len(self.instance_ids)

    def iterator(self, chunk_size):
        return iter(self.instance_ids)

    def __getitem__(self, item):
        return QuerySetStub(self.instance_ids[item])


def test_iterators_exclude_instances_completed_by_their_phase(caplog):
    command = Command.__new__(Command)
    command.dev_mode = False
    command.run_state = RunState(
        configuration=CONFIGURATION,
        started_at=datetime.datetime.now(),
        completed_ids={"validation": {1}, "recalculation": {2}},
    )
    command._base_qs = lambda: QuerySetStub([1, 2, 3])

    assert QuerySetStub([1, 2, 3])[1:].instance_ids == [2, 3]

    with caplog.at_level(logging.INFO):
        validation_total, validation_instances = command.iter_instances_for_validation()
        recalculation_total, recalculation_instances = (
            command.iter_instances_for_miscalculation()
        )

    assert validation_total == 2
    assert list(validation_instances) == [2, 3]
    assert recalculation_total == 2
    assert list(recalculation_instances) == [1, 3]
    assert "Skipping 1 instances already checked for validation" in caplog.messages
    assert "Skipping 1 instances already checked or fixed" in caplog.messages


def test_restored_results_populate_existing_excel_sheets(tmp_path):
    command = Command.__new__(Command)
    command.stats = defaultdict(lambda: defaultdict(list))
    command.run_state = RunState(
        configuration=CONFIGURATION,
        started_at=datetime.datetime.now(),
        results={
            ("recalculation", 1): [
                _stats_entry(1),
                _stats_entry(1, event="fixed"),
            ]
        },
    )
    command.analyze_only = None
    command.excel_out = tmp_path / "recalculation.xlsx"

    command._restore_stats()
    command.log_to_excel(
        command.run_state.started_at, committed_to_db=True, did_fixes=True
    )

    workbook = openpyxl.load_workbook(command.excel_out)
    assert workbook.sheetnames == ["Info", "Analysis", "Fixes", "Statistics"]
    assert workbook["Analysis"].max_row == 2
    assert workbook["Fixes"].max_row == 2
    assert workbook["Statistics"].max_row == 2
